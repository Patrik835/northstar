import csv
import hashlib
import io
import re
import unicodedata
import uuid
from dataclasses import dataclass
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from typing import Protocol

from openpyxl import load_workbook
from sqlalchemy import delete, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.encryption import CredentialCipher
from app.integrations.connectors.base import ConnectorPosition
from app.integrations.market_data import (
    BinanceCryptoPriceProvider,
    CryptoPriceError,
    EcbFxRateProvider,
    FxRateError,
    FxRateProvider,
)
from app.models.broker import BrokerConnection
from app.models.enums import AssetType, Broker, ConnectionStatus, TransactionType
from app.models.portfolio import PortfolioSnapshot, Position, Transaction
from app.repositories.connections import ConnectionRepository
from app.services.instrument_resolver import InstrumentResolver

MAX_BYBIT_FILE_BYTES = 20 * 1024 * 1024
MAX_BYBIT_FILES = 12
STABLECOINS = {"USDT", "USDC", "FDUSD", "TUSD", "DAI", "USDE"}
KNOWN_QUOTES = tuple(
    sorted(STABLECOINS | {"EUR", "USD", "BTC", "ETH"}, key=len, reverse=True)
)


class BybitImportError(ValueError):
    pass


class CryptoRates(Protocol):
    async def rates_to_eur(self, assets: set[str]) -> dict[str, Decimal]: ...


@dataclass(frozen=True, slots=True)
class BybitUpload:
    filename: str
    content: bytes


@dataclass(frozen=True, slots=True)
class ParsedBybitTransaction:
    external_id: str
    ticker: str
    transaction_type: TransactionType
    quantity: Decimal | None
    price: Decimal | None
    value: Decimal
    currency: str
    executed_at: datetime


@dataclass(frozen=True, slots=True)
class BalanceObservation:
    quantity: Decimal
    observed_at: datetime


@dataclass(frozen=True, slots=True)
class ParsedBybitFiles:
    rows_read: int
    transactions: list[ParsedBybitTransaction]
    balances: dict[str, BalanceObservation]
    warnings: list[str]


@dataclass(frozen=True, slots=True)
class BybitImportSummary:
    connection: BrokerConnection
    rows_read: int
    transactions_added: int
    duplicates_skipped: int
    positions_imported: int
    warnings: list[str]


ALIASES = {
    "datetime": {
        "transactiontime",
        "tradetime",
        "exectime",
        "executiontime",
        "ordertime",
        "updatetime",
        "createdtime",
        "createdat",
        "timestamp",
        "datetime",
        "date",
        "timeutc",
    },
    "type": {
        "type",
        "transactiontype",
        "transtype",
        "changetype",
        "recordtype",
        "operation",
    },
    "status": {"status", "orderstatus", "state"},
    "symbol": {"symbol", "pair", "tradingpair", "contract"},
    "side": {"side", "direction", "orderside"},
    "currency": {"currency", "coin", "asset", "token"},
    "quantity": {
        "qty",
        "quantity",
        "execqty",
        "executedqty",
        "filledqty",
        "filledquantity",
        "orderqty",
        "amount",
    },
    "price": {
        "price",
        "tradeprice",
        "execprice",
        "executionprice",
        "filledprice",
        "averagefilledprice",
        "avgfilledprice",
        "avgprice",
    },
    "value": {
        "value",
        "execvalue",
        "executedvalue",
        "filledvalue",
        "turnover",
        "quotequantity",
        "cashflow",
        "transactionamount",
    },
    "change": {"change", "balancechange", "walletchange"},
    "balance": {"cashbalance", "walletbalance", "balance", "endingbalance"},
    "fee": {"fee", "execfee", "tradingfee", "commission"},
    "fee_currency": {"feecurrency", "feecoin", "commissionasset"},
    "funding": {"funding", "fundingfee"},
    "trade_id": {"tradeid", "execid", "executionid"},
    "transaction_id": {"transactionid", "transid", "id", "recordid"},
    "order_id": {"orderid", "orderno", "ordernumber"},
    "tx_id": {"txid", "txhash", "transactionhash", "withdrawid"},
}
ALL_ALIASES = set().union(*ALIASES.values())


class BybitStatementParser:
    def parse(self, uploads: list[BybitUpload]) -> ParsedBybitFiles:
        if not uploads:
            raise BybitImportError("Select at least one Bybit CSV or XLSX export.")
        if len(uploads) > MAX_BYBIT_FILES:
            raise BybitImportError(f"Select at most {MAX_BYBIT_FILES} Bybit files at once.")

        transactions: list[ParsedBybitTransaction] = []
        balances: dict[str, BalanceObservation] = {}
        warnings: list[str] = []
        rows_read = 0
        recognized_tables = 0
        occurrences: dict[str, int] = {}
        for upload in uploads:
            if not upload.content:
                raise BybitImportError(f"{upload.filename} is empty.")
            if len(upload.content) > MAX_BYBIT_FILE_BYTES:
                raise BybitImportError(f"{upload.filename} is larger than the 20 MB limit.")
            suffix = upload.filename.casefold()
            if not suffix.endswith((".csv", ".xlsx")):
                raise BybitImportError("Bybit imports support CSV and XLSX files.")

            upload_recognized = False
            for table_name, headers, rows in _tables(upload):
                fields = _resolve_fields(headers)
                if not _recognized(fields):
                    continue
                upload_recognized = True
                recognized_tables += 1
                for row_number, row in enumerate(rows, start=2):
                    if not any(_text(value) for value in row.values()):
                        continue
                    rows_read += 1
                    try:
                        items, balance = self._parse_row(
                            table_name, row, fields, occurrences
                        )
                    except BybitImportError as exc:
                        raise BybitImportError(
                            f"{upload.filename}, row {row_number}: {exc}"
                        ) from exc
                    transactions.extend(items)
                    if balance is not None:
                        asset, observation = balance
                        existing = balances.get(asset)
                        if existing is None or observation.observed_at >= existing.observed_at:
                            balances[asset] = observation
            if not upload_recognized:
                warnings.append(f"{upload.filename} did not contain a recognized Bybit table.")

        if recognized_tables == 0:
            raise BybitImportError(
                "The selected files do not contain a recognized Bybit Transaction Log, "
                "Spot Trade History, or funding-history table."
            )
        deduplicated = list({item.external_id: item for item in transactions}.values())
        if not deduplicated and not balances:
            raise BybitImportError("The selected Bybit exports contain no supported activity.")
        if not balances:
            warnings.append(
                "No Account Change Details cash balances were found; current holdings "
                "will be reconstructed from the imported activity."
            )
        return ParsedBybitFiles(rows_read, deduplicated, balances, _unique(warnings))

    def _parse_row(
        self,
        table_name: str,
        row: dict[str, object],
        fields: dict[str, str],
        occurrences: dict[str, int],
    ) -> tuple[
        list[ParsedBybitTransaction], tuple[str, BalanceObservation] | None
    ]:
        status = _normalize(_field(row, fields, "status"))
        if any(token in status for token in ("cancel", "reject", "failed", "pending")):
            return [], None
        executed_at = _datetime(_field(row, fields, "datetime"))
        raw_type = _field(row, fields, "type")
        side = _normalize(_field(row, fields, "side"))
        symbol = _symbol(_field(row, fields, "symbol"))
        asset = _asset(_field(row, fields, "currency"))
        balance_value = _decimal(_field(row, fields, "balance"))
        balance = (
            (asset, BalanceObservation(balance_value, executed_at))
            if asset and balance_value is not None
            else None
        )

        if symbol and side in {"buy", "sell"}:
            return self._trade(row, fields, symbol, side, executed_at, occurrences), balance

        transaction_type = _flow_type(f"{raw_type} {table_name}", row, fields)
        if transaction_type is None:
            return [], balance
        quantity = _decimal(
            _field(row, fields, "change") or _field(row, fields, "quantity")
        )
        if not asset or quantity is None or quantity == 0:
            return [], balance
        identifier = _identifier(row, fields) or _fingerprint(row, occurrences)
        prefix = {
            TransactionType.DEPOSIT: "deposit",
            TransactionType.WITHDRAWAL: "withdrawal",
            TransactionType.DIVIDEND: "income",
            TransactionType.FEE: "fee",
        }[transaction_type]
        item = ParsedBybitTransaction(
            external_id=f"{prefix}:{identifier}"[:160],
            ticker=asset,
            transaction_type=transaction_type,
            quantity=abs(quantity),
            price=None,
            value=abs(quantity),
            currency=asset,
            executed_at=executed_at,
        )
        return [item], balance

    def _trade(
        self,
        row: dict[str, object],
        fields: dict[str, str],
        symbol: str,
        side: str,
        executed_at: datetime,
        occurrences: dict[str, int],
    ) -> list[ParsedBybitTransaction]:
        base, quote = _pair(symbol, _asset(_field(row, fields, "currency")))
        quantity = _decimal(_field(row, fields, "quantity"))
        price = _decimal(_field(row, fields, "price"))
        value = _decimal(_field(row, fields, "value"))
        if quantity is None or quantity == 0:
            raise BybitImportError("a completed Spot trade has no filled quantity")
        quantity = abs(quantity)
        if value is None and price is not None:
            value = quantity * abs(price)
        if price is None and value is not None:
            price = abs(value) / quantity
        if price is None or value is None:
            raise BybitImportError("a completed Spot trade has no usable filled price")
        identifier = _identifier(row, fields) or _fingerprint(row, occurrences)
        items = [
            ParsedBybitTransaction(
                external_id=f"trade:{symbol}:{identifier}"[:160],
                ticker=base,
                transaction_type=(
                    TransactionType.BUY if side == "buy" else TransactionType.SELL
                ),
                quantity=quantity,
                price=abs(price),
                value=abs(value),
                currency=quote,
                executed_at=executed_at,
            )
        ]
        fee = _decimal(_field(row, fields, "fee"))
        if fee is not None and fee != 0:
            fee_asset = _asset(_field(row, fields, "fee_currency")) or quote
            items.append(
                ParsedBybitTransaction(
                    external_id=f"trade-fee:{symbol}:{identifier}"[:160],
                    ticker=fee_asset,
                    transaction_type=TransactionType.FEE,
                    quantity=abs(fee),
                    price=None,
                    value=abs(fee),
                    currency=fee_asset,
                    executed_at=executed_at,
                )
            )
        return items


class BybitStatementImportService:
    def __init__(
        self,
        db: AsyncSession,
        cipher: CredentialCipher,
        prices: CryptoRates | None = None,
        fx_rates: FxRateProvider | None = None,
    ) -> None:
        self.db = db
        self.cipher = cipher
        self.prices = prices or BinanceCryptoPriceProvider()
        self.fx_rates = fx_rates or EcbFxRateProvider(db)

    async def import_files(
        self, user_id: uuid.UUID, uploads: list[BybitUpload]
    ) -> BybitImportSummary:
        parsed = BybitStatementParser().parse(uploads)
        connection = await ConnectionRepository(self.db).by_broker(user_id, Broker.BYBIT)
        if connection is None:
            connection = BrokerConnection(
                user_id=user_id,
                broker=Broker.BYBIT,
                encrypted_credentials=self.cipher.encrypt({"source": "statement"}),
                credential_hint="File import",
                status=ConnectionStatus.PENDING,
            )
            self.db.add(connection)
            await self.db.flush()
        else:
            # An earlier API connection may still contain a key. The first statement
            # import deliberately replaces it so no obsolete Bybit secret is retained.
            connection.encrypted_credentials = self.cipher.encrypt({"source": "statement"})
            connection.credential_hint = "File import"

        existing_ids = set(
            await self.db.scalars(
                select(Transaction.external_id).where(
                    Transaction.broker_connection_id == connection.id
                )
            )
        )
        assets = {
            item.currency for item in parsed.transactions
        } | set(parsed.balances)
        try:
            rates = await self.prices.rates_to_eur(assets)
        except CryptoPriceError:
            rates = {}

        warnings = list(parsed.warnings)
        converted: list[tuple[ParsedBybitTransaction, Decimal | None]] = []
        for item in parsed.transactions:
            converted.append((item, await self._value_eur(item, rates, warnings)))
        new_items = [item for item in converted if item[0].external_id not in existing_ids]
        self.db.add_all(
            [
                Transaction(
                    broker_connection_id=connection.id,
                    external_id=item.external_id,
                    ticker=item.ticker,
                    transaction_type=item.transaction_type,
                    quantity=item.quantity,
                    price=item.price,
                    value=item.value,
                    value_eur=value_eur,
                    currency=item.currency,
                    executed_at=item.executed_at,
                )
                for item, value_eur in new_items
            ]
        )
        await self.db.flush()

        all_transactions = list(
            await self.db.scalars(
                select(Transaction)
                .where(Transaction.broker_connection_id == connection.id)
                .order_by(Transaction.executed_at, Transaction.external_id)
            )
        )
        states = _reconstruct(all_transactions)
        quantities = {
            asset: observation.quantity
            for asset, observation in parsed.balances.items()
            if observation.quantity > 0
        }
        for asset, state in states.items():
            if asset not in parsed.balances and state.quantity > 0:
                quantities[asset] = state.quantity

        await self.db.execute(
            delete(Position).where(Position.broker_connection_id == connection.id)
        )
        resolver = InstrumentResolver(self.db)
        positions: list[Position] = []
        total_value_eur = Decimal(0)
        valuation_at = datetime.now(timezone.utc)
        valuation_limited = not bool(parsed.balances)
        for ticker, quantity in sorted(quantities.items()):
            state = states.get(ticker, _HoldingState())
            rate = rates.get(ticker) or state.last_price_eur
            if rate is None:
                warnings.append(f"{ticker} was skipped because no EUR market price was available.")
                valuation_limited = True
                continue
            if ticker not in rates:
                warnings.append(f"{ticker} uses its last imported trade price.")
                valuation_limited = True
            value_eur = quantity * rate
            total_value_eur += value_eur
            quantity_matches = abs(state.quantity - quantity) <= Decimal("0.00000001")
            average_price = (
                state.cost_eur / quantity
                if quantity and state.cost_complete and quantity_matches
                else None
            )
            connector_position = ConnectorPosition(
                instrument_id=f"BYBIT:{ticker}",
                ticker=ticker,
                name=ticker,
                asset_type=AssetType.CRYPTO,
                quantity=quantity,
                average_price=average_price,
                current_value=value_eur,
                currency="EUR",
                canonical_symbol=ticker,
            )
            instrument = await resolver.resolve(Broker.BYBIT, connector_position)
            positions.append(
                Position(
                    broker_connection_id=connection.id,
                    instrument_id=connector_position.instrument_id,
                    canonical_instrument_id=instrument.id,
                    ticker=ticker,
                    name=ticker,
                    asset_type=AssetType.CRYPTO,
                    quantity=quantity,
                    average_price=average_price,
                    current_value=value_eur,
                    currency="EUR",
                    current_value_eur=value_eur,
                    valued_at=valuation_at,
                    valuation_source="market" if ticker in rates else "last_trade",
                    is_estimated=ticker not in rates or not bool(parsed.balances),
                )
            )
        self.db.add_all(positions)

        snapshot = await self.db.scalar(
            select(PortfolioSnapshot).where(
                PortfolioSnapshot.broker_connection_id == connection.id,
                PortfolioSnapshot.snapshot_date == date.today(),
            )
        )
        if snapshot is None:
            self.db.add(
                PortfolioSnapshot(
                    broker_connection_id=connection.id,
                    snapshot_date=date.today(),
                    total_value=total_value_eur,
                    total_value_eur=total_value_eur,
                    currency="EUR",
                )
            )
        else:
            snapshot.total_value = total_value_eur
            snapshot.total_value_eur = total_value_eur
            snapshot.currency = "EUR"
            snapshot.reported_pnl = None
            snapshot.reported_pnl_eur = None

        completed_at = datetime.now(timezone.utc)
        connection.status = (
            ConnectionStatus.LIMITED if valuation_limited else ConnectionStatus.ACTIVE
        )
        connection.last_error = None
        connection.last_sync_attempt_at = completed_at
        connection.last_successful_sync_at = completed_at
        connection.last_synced_at = completed_at
        connection.reconciliation_difference_percent = None
        connection.reconciliation_checked_at = None
        connection.reconciliation_warning = None
        await self.db.commit()
        await self.db.refresh(connection)
        return BybitImportSummary(
            connection=connection,
            rows_read=parsed.rows_read,
            transactions_added=len(new_items),
            duplicates_skipped=len(parsed.transactions) - len(new_items),
            positions_imported=len(positions),
            warnings=_unique(warnings),
        )

    async def _value_eur(
        self,
        item: ParsedBybitTransaction,
        rates: dict[str, Decimal],
        warnings: list[str],
    ) -> Decimal | None:
        currency = item.currency
        try:
            if currency in STABLECOINS:
                return await self.fx_rates.convert_to_eur(
                    item.value, "USD", item.executed_at.date()
                )
            if currency in {"EUR", "USD", "GBP", "CHF", "CZK", "PLN"}:
                return await self.fx_rates.convert_to_eur(
                    item.value, currency, item.executed_at.date()
                )
        except FxRateError:
            pass
        if currency in rates:
            if item.transaction_type not in {TransactionType.BUY, TransactionType.SELL}:
                warnings.append(
                    "Crypto-denominated transfers, income, and fees use import-time EUR prices."
                )
            return item.value * rates[currency]
        warnings.append(f"No EUR conversion was available for {currency} activity.")
        return None


@dataclass(slots=True)
class _HoldingState:
    quantity: Decimal = Decimal(0)
    cost_eur: Decimal = Decimal(0)
    last_price_eur: Decimal | None = None
    cost_complete: bool = True


def _reconstruct(transactions: list[Transaction]) -> dict[str, _HoldingState]:
    states: dict[str, _HoldingState] = {}
    for item in transactions:
        ticker = item.ticker.upper()
        quantity = abs(item.quantity or Decimal(0))
        if not ticker or not quantity:
            continue
        state = states.setdefault(ticker, _HoldingState())
        if item.transaction_type is TransactionType.BUY:
            state.quantity += quantity
            if item.value_eur is None:
                state.cost_complete = False
            else:
                state.cost_eur += item.value_eur
                state.last_price_eur = item.value_eur / quantity
        elif item.transaction_type is TransactionType.SELL:
            if quantity > state.quantity:
                state.cost_complete = False
            removed = min(quantity, state.quantity)
            average = state.cost_eur / state.quantity if state.quantity else Decimal(0)
            state.quantity -= removed
            state.cost_eur = max(Decimal(0), state.cost_eur - average * removed)
            if item.value_eur is not None:
                state.last_price_eur = item.value_eur / quantity
        elif item.transaction_type is TransactionType.DEPOSIT:
            state.quantity += quantity
            state.cost_complete = False
        elif item.transaction_type is TransactionType.WITHDRAWAL:
            removed = min(quantity, state.quantity)
            average = state.cost_eur / state.quantity if state.quantity else Decimal(0)
            state.quantity -= removed
            state.cost_eur = max(Decimal(0), state.cost_eur - average * removed)
            if quantity > removed:
                state.cost_complete = False
        elif item.transaction_type is TransactionType.DIVIDEND:
            state.quantity += quantity
        elif item.transaction_type is TransactionType.FEE:
            removed = min(quantity, state.quantity)
            average = state.cost_eur / state.quantity if state.quantity else Decimal(0)
            state.quantity -= removed
            state.cost_eur = max(Decimal(0), state.cost_eur - average * removed)
    return states


def _tables(upload: BybitUpload):
    if upload.filename.casefold().endswith(".xlsx"):
        try:
            workbook = load_workbook(io.BytesIO(upload.content), read_only=True, data_only=True)
        except Exception as exc:
            raise BybitImportError(f"{upload.filename} is not a readable XLSX file.") from exc
        for sheet in workbook.worksheets:
            matrix = [list(row) for row in sheet.iter_rows(values_only=True)]
            table = _matrix_table(matrix)
            if table is not None:
                headers, rows = table
                yield f"{upload.filename} {sheet.title}", headers, rows
        return
    try:
        text = _decode(upload.content)
        dialect = csv.Sniffer().sniff(text[:16384], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    matrix = list(csv.reader(io.StringIO(text), dialect=dialect))
    table = _matrix_table(matrix)
    if table is not None:
        headers, rows = table
        yield upload.filename, headers, rows


def _matrix_table(matrix: list[list[object]]):
    best_index = -1
    best_score = 0
    for index, row in enumerate(matrix[:30]):
        score = sum(1 for value in row if _normalize(value) in ALL_ALIASES)
        if score > best_score:
            best_index, best_score = index, score
    if best_index < 0 or best_score < 2:
        return None
    headers = [_text(value) or f"column_{index}" for index, value in enumerate(matrix[best_index])]
    rows = [dict(zip(headers, row, strict=False)) for row in matrix[best_index + 1 :]]
    return headers, rows


def _resolve_fields(headers: list[str]) -> dict[str, str]:
    normalized = {_normalize(header): header for header in headers}
    return {
        field: normalized[alias]
        for field, aliases in ALIASES.items()
        if (alias := next((value for value in aliases if value in normalized), None))
    }


def _recognized(fields: dict[str, str]) -> bool:
    return "datetime" in fields and (
        {"symbol", "side", "quantity"}.issubset(fields)
        or "balance" in fields
        or ("type" in fields and ("currency" in fields or "symbol" in fields))
    )


def _field(row: dict[str, object], fields: dict[str, str], name: str) -> str:
    return _text(row.get(fields.get(name, "")))


def _identifier(row: dict[str, object], fields: dict[str, str]) -> str:
    for field in ("trade_id", "transaction_id", "order_id", "tx_id"):
        if value := _field(row, fields, field):
            return re.sub(r"\s+", "", value)[:120]
    return ""


def _fingerprint(row: dict[str, object], occurrences: dict[str, int]) -> str:
    digest = hashlib.sha256(
        "|".join(
            f"{_normalize(key)}={_text(value)}" for key, value in sorted(row.items())
        ).encode()
    ).hexdigest()[:32]
    occurrences[digest] = occurrences.get(digest, 0) + 1
    return f"{digest}:{occurrences[digest]}"


def _flow_type(
    value: str, row: dict[str, object], fields: dict[str, str]
) -> TransactionType | None:
    normalized = _normalize(value)
    change = _decimal(_field(row, fields, "change") or _field(row, fields, "quantity"))
    if any(token in normalized for token in ("deposit", "transferin", "fiatdeposit")):
        return TransactionType.DEPOSIT
    if any(token in normalized for token in ("withdraw", "transferout", "fiatwithdraw")):
        return TransactionType.WITHDRAWAL
    if any(
        token in normalized
        for token in ("yield", "interest", "reward", "airdrop", "earn", "cashback", "bonus")
    ):
        return TransactionType.DIVIDEND
    if any(token in normalized for token in ("fee", "commission", "funding")):
        return TransactionType.FEE
    if "transfer" in normalized and change is not None:
        return TransactionType.DEPOSIT if change > 0 else TransactionType.WITHDRAWAL
    return None


def _pair(symbol: str, currency_hint: str) -> tuple[str, str]:
    normalized = _symbol(symbol)
    for separator in ("/", "-", "_"):
        if separator in symbol:
            base, quote = (_asset(value) for value in symbol.split(separator, 1))
            if base and quote:
                return base, quote
    hints = tuple(value for value in (currency_hint, *KNOWN_QUOTES) if value)
    for quote in hints:
        if normalized.endswith(quote) and len(normalized) > len(quote):
            return normalized[: -len(quote)], quote
    raise BybitImportError(f"'{symbol}' is not a recognized Spot pair")


def _datetime(value: str) -> datetime:
    raw = value.strip()
    if not raw:
        raise BybitImportError("a row has no transaction time")
    if raw.isdigit():
        number = int(raw)
        if number > 10_000_000_000:
            number //= 1000
        try:
            return datetime.fromtimestamp(number, tz=timezone.utc)
        except (OverflowError, OSError, ValueError):
            pass
    normalized = re.sub(r"\s*\(?UTC[+-]?\d*\)?\s*$", "", raw, flags=re.IGNORECASE)
    normalized = normalized.replace("Z", "+00:00")
    try:
        parsed = datetime.fromisoformat(normalized)
    except ValueError:
        parsed = None
        for pattern in (
            "%Y-%m-%d %H:%M:%S",
            "%Y-%m-%d %H:%M",
            "%d/%m/%Y %H:%M:%S",
            "%d.%m.%Y %H:%M:%S",
            "%m/%d/%Y %H:%M:%S",
        ):
            try:
                parsed = datetime.strptime(normalized, pattern)
                break
            except ValueError:
                continue
        if parsed is None:
            raise BybitImportError(f"'{raw}' is not a recognized transaction time") from None
    return parsed if parsed.tzinfo else parsed.replace(tzinfo=timezone.utc)


def _decimal(value: str) -> Decimal | None:
    cleaned = value.strip()
    if not cleaned or cleaned.upper() in {"-", "—", "N/A", "NULL", "NONE"}:
        return None
    negative = cleaned.startswith("(") and cleaned.endswith(")")
    cleaned = re.sub(r"[^0-9,.-]", "", cleaned.strip("()"))
    if "," in cleaned and "." in cleaned:
        cleaned = cleaned.replace(",", "")
    elif "," in cleaned:
        cleaned = cleaned.replace(",", ".")
    try:
        result = Decimal(cleaned)
    except InvalidOperation as exc:
        raise BybitImportError(f"'{value}' is not a valid number") from exc
    return -result if negative else result


def _decode(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise BybitImportError("The CSV text encoding is not supported.")


def _normalize(value: object) -> str:
    normalized = unicodedata.normalize("NFKD", _text(value))
    ascii_value = "".join(char for char in normalized if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]", "", ascii_value.casefold())


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat()
    return str(value).strip()


def _asset(value: str) -> str:
    return re.sub(r"[^A-Z0-9]", "", value.upper())[:32]


def _symbol(value: str) -> str:
    return re.sub(r"[^A-Z0-9/_-]", "", value.upper())[:40]


def _unique(items: list[str]) -> list[str]:
    return list(dict.fromkeys(item for item in items if item))
